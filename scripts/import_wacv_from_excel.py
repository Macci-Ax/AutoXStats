#!/usr/bin/env python3
"""
Import WACV championship data from Meisterschaftswertung_2025_WACV.xlsx
Uses the same structure as DRCV (physical_events, championship_events, class_events)
"""

import sqlite3
import openpyxl
import os
import re
import uuid

DB_PATH = 'autox.db'
EXCEL_PATH = os.path.join('pdf', 'Meisterschaftswertung_2025_WACV.xlsx')

# Event configuration matching the Excel columns
WACV_EVENTS = [
    {'name': 'AC Waldorf', 'date': '2025-04-06', 'location': 'Waldorf'},
    {'name': 'MSC Extertal', 'date': '2025-04-27', 'location': 'Extertal'},
    {'name': 'RSC Düdinghausen', 'date': '2025-05-18', 'location': 'Düdinghausen'},
    {'name': 'RSG Aartal', 'date': '2025-06-08', 'location': 'Eppe'},
    {'name': 'MSC Linsburg', 'date': '2025-06-22', 'location': 'Linsburg'},
    {'name': 'MSF Lichtenau', 'date': '2025-07-13', 'location': 'Lichtenau'},
    {'name': 'RT Oberschledorn', 'date': '2025-08-03', 'location': 'Oberschledorn'},
    {'name': 'MSC Crazy Horses', 'date': '2025-08-24', 'location': 'Crazy Horses'},
    {'name': 'MSC Hesborn', 'date': '2025-09-14', 'location': 'Hesborn'}
]

def get_db_connection():
    return sqlite3.connect(DB_PATH)

def format_driver_name(raw_name):
    """
    Format driver name:
    1. Take only first name before '/' (for team entries)
    2. Convert 'Nachname, Vorname' to 'Vorname Nachname'
    """
    if not raw_name:
        return ''
    
    # Step 1: Take first name only (before any '/')
    first_name = raw_name.split('/')[0].strip()
    
    # Step 2: Convert "Nachname, Vorname" to "Vorname Nachname"
    if ',' in first_name:
        parts = first_name.split(',', 1)  # Split only on first comma
        nachname = parts[0].strip()
        vorname = parts[1].strip() if len(parts) > 1 else ''
        if vorname:
            return f"{vorname} {nachname}"
        return nachname
    
    return first_name

def sheet_name_to_class_id(sheet_name):
    """Convert sheet name to class_id."""
    lower = sheet_name.lower().strip()
    
    if 'jugend' in lower and 'langstrecke' in lower:
        return 'w_jugend_lang'
    elif 'jugend' in lower:
        match = re.search(r'klasse\s*(\d+)', lower)
        if match:
            return f'w_jugend_k{match.group(1)}'
        return 'w_jugend'
    elif 'langstrecke' in lower:
        return 'w_lang'
    elif 'endlauf' in lower and 'touren' in lower:
        return 'w_endlauf_touren'
    elif 'endlauf' in lower and 'spezial' in lower:
        return 'w_endlauf_spezial'
    elif 'endlauf' in lower:
        return 'w_endlauf'
    else:
        match = re.search(r'klasse\s*(\d+)', lower)
        if match:
            return f'w_k{match.group(1)}'
    
    clean = lower.replace(' ', '_')
    return f'w_{clean}'

def cleanup_wacv_data(cursor):
    """Remove all existing WACV data."""
    print("Cleaning old WACV data...")
    
    # Delete in correct order to respect foreign keys
    cursor.execute("DELETE FROM race_results WHERE class_id LIKE 'w_%'")
    cursor.execute("DELETE FROM driver_participations WHERE class_id LIKE 'w_%'")
    cursor.execute("DELETE FROM class_events WHERE class_id LIKE 'w_%'")
    cursor.execute("DELETE FROM drivers WHERE id LIKE 'w_%'")
    cursor.execute("DELETE FROM classes WHERE id LIKE 'w_%'")
    cursor.execute("DELETE FROM championship_events WHERE championship_id = 'WACV'")
    cursor.execute("DELETE FROM physical_events WHERE id LIKE 'pe_wacv_%'")
    
    print("  Old WACV data removed.")

def setup_wacv_infrastructure(cursor):
    """Create WACV championship, events, and classes structure (like DRCV)."""
    print("\nSetting up WACV infrastructure...")
    
    # 1. Ensure WACV championship exists
    cursor.execute("""
        INSERT OR IGNORE INTO championships (id, name, year) 
        VALUES ('WACV', 'Westdeutscher Auto Cross Verband', 2025)
    """)
    
    # 2. Create physical_events and championship_events for each WACV event
    event_ids = []
    for i, event in enumerate(WACV_EVENTS, start=1):
        pe_id = f"pe_wacv_{i}"
        ce_id = f"ce_wacv_{i}"
        
        # Physical Event
        cursor.execute("""
            INSERT OR REPLACE INTO physical_events (id, title, start_date, location, status)
            VALUES (?, ?, ?, ?, 'finished')
        """, (pe_id, event['name'], event['date'], event['location']))
        
        # Championship Event
        cursor.execute("""
            INSERT OR REPLACE INTO championship_events (id, physical_event_id, championship_id, has_results)
            VALUES (?, ?, 'WACV', 1)
        """, (ce_id, pe_id))
        
        event_ids.append(ce_id)
        print(f"  Created event: {event['name']} -> {ce_id}")
    
    return event_ids

def find_or_create_class(cursor, sheet_name, event_ids):
    """Find or create a class and link it to all events via class_events."""
    class_id = sheet_name_to_class_id(sheet_name)
    
    # Create class if not exists
    cursor.execute("SELECT id FROM classes WHERE id = ?", (class_id,))
    if not cursor.fetchone():
        print(f"  Creating class: {sheet_name} -> {class_id}")
        cursor.execute("""
            INSERT INTO classes (id, championship_id, name) 
            VALUES (?, 'WACV', ?)
        """, (class_id, sheet_name))
    
    # Create class_events for each event (discipline = 'klasse' for regular, 'langstrecke' for Langstrecke)
    discipline = 'langstrecke' if 'langstrecke' in sheet_name.lower() else 'klasse'
    
    for ce_id in event_ids:
        cursor.execute("""
            INSERT OR REPLACE INTO class_events (class_id, event_id, discipline)
            VALUES (?, ?, ?)
        """, (class_id, ce_id, discipline))
    
    return class_id

def find_or_create_driver(cursor, name, team, number, class_id):
    """Find or create a driver."""
    if not name:
        return None
    
    # Check if driver exists by name (case-insensitive)
    cursor.execute("SELECT id FROM drivers WHERE name = ? COLLATE NOCASE", (name,))
    row = cursor.fetchone()
    if row:
        return row[0]
    
    # Create new driver
    new_id = f"w_{number}_{uuid.uuid4().hex[:4]}"
    cursor.execute("""
        INSERT INTO drivers (id, name, team, car, start_number, current_class_id)
        VALUES (?, ?, ?, '', ?, ?)
    """, (new_id, name, team or '', number, class_id))
    return new_id

def process_sheet(cursor, sheet, sheet_name, event_ids):
    """Process a single sheet (class)."""
    print(f"\n=== Processing Sheet: {sheet_name} ===")
    
    class_id = find_or_create_class(cursor, sheet_name, event_ids)
    
    # Read header to get column mapping
    header_row = [cell.value for cell in list(sheet.iter_rows(min_row=1, max_row=1))[0]]
    
    # Build event column mapping (match header to event_ids)
    event_columns = {}
    for col_idx, header in enumerate(header_row):
        if header and col_idx >= 6:  # Event columns start after Position, Startnr, Team, Fahrer, Club, Punkte
            for i, event in enumerate(WACV_EVENTS):
                if event['name'] in str(header):
                    event_columns[col_idx] = event_ids[i]
                    break
    
    print(f"  Event columns mapped: {len(event_columns)}")
    
    drivers_processed = 0
    
    # Process data rows (starting from row 2)
    for row_idx, row in enumerate(sheet.iter_rows(min_row=2), start=2):
        values = [cell.value for cell in row]
        
        if not values[0] or not isinstance(values[0], (int, float)):
            continue
        
        try:
            rank = int(values[0])
            number = int(values[1]) if values[1] else 0
            team = str(values[2]) if values[2] else ''
            driver_name = str(values[3]) if values[3] else ''
            club = str(values[4]) if values[4] else ''
            total_points = int(values[5]) if values[5] else 0
            
            if not driver_name:
                continue
            
            # Format driver name: "Nachname, Vorname" -> "Vorname Nachname", only first before "/"
            driver_name = format_driver_name(driver_name)
            
            # Find or create driver
            driver_id = find_or_create_driver(cursor, driver_name, team, number, class_id)
            if not driver_id:
                continue
            
            # Insert/Update participation
            cursor.execute("""
                INSERT OR REPLACE INTO driver_participations (driver_id, class_id, points, rank, wins, podiums)
                VALUES (?, ?, ?, ?, 0, 0)
            """, (driver_id, class_id, total_points, rank))
            
            # Insert race results for each event
            for col_idx, ce_id in event_columns.items():
                if col_idx < len(values) and values[col_idx]:
                    try:
                        pts = int(values[col_idx])
                        if pts > 0:
                            # Use championship_event_id (like DRCV does)
                            res_id = f"w_res_{driver_id}_{ce_id}"
                            cursor.execute("""
                                INSERT OR REPLACE INTO race_results 
                                (id, championship_event_id, driver_id, class_id, championship_points, rank)
                                VALUES (?, ?, ?, ?, ?, ?)
                            """, (res_id, ce_id, driver_id, class_id, pts, rank))
                    except (ValueError, TypeError):
                        pass
            
            drivers_processed += 1
            if drivers_processed <= 3:
                print(f"   #{rank} {driver_name} (Pts: {total_points})")
                
        except Exception as e:
            print(f"  Error processing row {row_idx}: {e}")
            continue
    
    print(f"  Total drivers processed: {drivers_processed}")
    return drivers_processed

def main():
    print(f"Importing WACV data from: {EXCEL_PATH}")
    print("Using DRCV-compatible structure (physical_events, championship_events, class_events)")
    
    if not os.path.exists(EXCEL_PATH):
        print(f"ERROR: File not found: {EXCEL_PATH}")
        return
    
    # Load workbook
    wb = openpyxl.load_workbook(EXCEL_PATH)
    print(f"Found {len(wb.sheetnames)} sheets: {wb.sheetnames}")
    
    # Connect to database
    conn = get_db_connection()
    cursor = conn.cursor()
    
    try:
        # 1. Clean up old WACV data
        cleanup_wacv_data(cursor)
        
        # 2. Setup infrastructure (events, etc.)
        event_ids = setup_wacv_infrastructure(cursor)
        
        # 3. Process each sheet
        total_drivers = 0
        for sheet_name in wb.sheetnames:
            sheet = wb[sheet_name]
            drivers = process_sheet(cursor, sheet, sheet_name, event_ids)
            total_drivers += drivers
        
        conn.commit()
        print(f"\n=== Import Complete ===")
        print(f"Total drivers imported: {total_drivers}")
        print(f"Sheets processed: {len(wb.sheetnames)}")
        print(f"WACV Events created: {len(event_ids)}")
        
    except Exception as e:
        print(f"ERROR: {e}")
        conn.rollback()
        raise
    finally:
        conn.close()

if __name__ == '__main__':
    main()
